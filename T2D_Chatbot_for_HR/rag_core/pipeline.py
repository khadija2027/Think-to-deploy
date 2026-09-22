"""Document stages shared by Airflow and regression tests.

Rebuild from the current corpus when its fingerprint changes. Failed runs never
replace the last good snapshot or mark input files as processed.
"""
import hashlib
import json
import os
import re
import subprocess
import uuid
from datetime import datetime, timezone
from pathlib import Path

from .storage import PROJECT_ROOT, current_manifest, index_directory, read_json, write_json

SUPPORTED = {".pdf", ".docx", ".doc", ".xlsx", ".xls", ".txt", ".md"}
CHUNK_SIZE = 1000
CHUNK_OVERLAP = 200
MAX_CHUNKS = 5000
PIPELINE_VERSION = 3


def discover(run_id):
    dataset = Path(os.environ.get("DATASET_DIR", PROJECT_ROOT / "dataset"))
    if not dataset.is_dir():
        raise ValueError(f"Dataset directory does not exist: {dataset}")
    model = os.environ.get("EMBEDDING_MODEL", "paraphrase-multilingual-MiniLM-L12-v2")
    files = []
    for path in sorted(dataset.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED:
            continue
        if any(part.startswith(".") for part in path.relative_to(dataset).parts):
            continue
        if not 0 < path.stat().st_size <= 50 * 1024 * 1024:
            raise ValueError(f"Empty or oversized document: {path.name}")
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        files.append({"path": str(path), "source": path.relative_to(dataset).as_posix(), "sha256": digest})
    signature = {"files": [(f["source"], f["sha256"]) for f in files],
                 "model": model, "chunk_size": CHUNK_SIZE, "overlap": CHUNK_OVERLAP,
                 "pipeline_version": PIPELINE_VERSION}
    fingerprint = hashlib.sha256(json.dumps(signature, sort_keys=True).encode()).hexdigest()
    previous = current_manifest()
    if previous and previous.get("fingerprint") == fingerprint:
        return None
    if not files and previous is None:
        return None
    work_root = Path(os.environ.get("PIPELINE_WORK_DIR", PROJECT_ROOT / "pipeline_work"))
    work = work_root / hashlib.sha256(run_id.encode()).hexdigest()[:24]
    work.mkdir(parents=True, exist_ok=True)
    write_json(work / "discovery.json", {"files": files, "fingerprint": fingerprint, "model": model})
    return str(work)


def parse_document(path):
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8-sig")
    if suffix == ".docx":
        from docx import Document
        document = Document(path)
        paragraphs = [p.text for p in document.paragraphs]
        paragraphs.extend(" | ".join(c.text for c in row.cells)
                          for table in document.tables for row in table.rows)
        return "\n\n".join(paragraphs)
    if suffix == ".doc":
        return subprocess.run(["antiword", str(path)], check=True, capture_output=True,
                              timeout=120).stdout.decode("utf-8", errors="replace")
    if suffix in {".xlsx", ".xls"}:
        import pandas as pd
        sheets = pd.read_excel(path, sheet_name=None, header=None)
        return "\n\n".join(f"Sheet: {name}\n{frame.fillna('').to_csv(index=False, header=False)}"
                           for name, frame in sheets.items())
    if suffix == ".pdf":
        import pdfplumber
        from pypdf import PdfReader
        pages = []
        fallback_reader = None
        with pdfplumber.open(path) as pdf:
            for number, page in enumerate(pdf.pages, 1):
                text = page.extract_text() or ""
                if not text.strip():
                    if fallback_reader is None:
                        fallback_reader = PdfReader(path)
                    text = fallback_reader.pages[number - 1].extract_text() or ""
                if not text.strip():
                    from pdf2image import convert_from_path
                    import pytesseract
                    images = convert_from_path(path, dpi=200, first_page=number,
                                               last_page=number, timeout=120)
                    text = "\n".join(pytesseract.image_to_string(image, lang="fra+eng", timeout=120)
                                     for image in images)
                for table in page.extract_tables():
                    text += "\n" + "\n".join(" | ".join(str(c or "") for c in row) for row in table)
                if text.strip():
                    pages.append(f"[Page {number}]\n{text}")
        return "\n\n".join(pages)
    raise ValueError(f"Unsupported format: {suffix}")


def parse_documents(work):
    discovery = read_json(Path(work) / "discovery.json")
    documents = []
    for item in discovery["files"]:
        path = Path(item["path"])
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        if before != item["sha256"]:
            raise ValueError(f"Document changed during ingestion: {item['source']}")
        text = parse_document(path)
        if not text.strip():
            raise ValueError(f"No text could be extracted: {item['source']}")
        if hashlib.sha256(path.read_bytes()).hexdigest() != before:
            raise ValueError(f"Document changed during parsing: {item['source']}")
        documents.append({"source": item["source"], "sha256": before, "text": text})
    write_json(Path(work) / "parsed.json", documents)
    return work


def anonymize_text(text):
    # Best-effort masking, not a claim of complete anonymization.
    patterns = [
        (r"\b[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}\b", "[EMAIL]"),
        (r"(?<!\w)(?:\+33|0)[1-9](?:[.\s-]?\d{2}){4}(?!\d)", "[PHONE]"),
        (r"\b(?:MAT|EMP|ID)[ -]?\d{4,8}\b", "[EMPLOYEE_ID]"),
        (r"\b[A-Z]{2}\d{2}(?:[ ]?[A-Z0-9]){11,30}\b", "[IBAN]"),
        (r"\b[12]\d{14}\b", "[SOCIAL_SECURITY]"),
        (r"\b(?:M\.|Mme|Monsieur|Madame)\s+[A-ZÀ-ÖØ-Þ][a-zà-öø-ÿ]+(?:\s+[A-ZÀ-ÖØ-Þ][a-zà-öø-ÿ]+)+", "[NAME]"),
    ]
    for pattern, replacement in patterns:
        text = re.sub(pattern, replacement, text)
    return text


def anonymize_documents(work):
    documents = read_json(Path(work) / "parsed.json")
    for document in documents:
        document["text"] = anonymize_text(document["text"])
    write_json(Path(work) / "anonymized.json", documents)
    return work


def chunk_text(text, size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
    if size <= 0 or not 0 <= overlap < size:
        raise ValueError("Chunk overlap must be smaller than chunk size")
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        piece = text[start:end].strip()
        if piece:
            chunks.append(piece)
        if end == len(text):
            break
        start = end - overlap
    return chunks


def chunk_documents(work):
    documents = read_json(Path(work) / "anonymized.json")
    chunks = []
    for document in documents:
        for number, text in enumerate(chunk_text(document["text"])):
            chunks.append({"text": text, "metadata": {
                "source_file": document["source"], "chunk_id": number,
                "sha256": document["sha256"], "pii_masking": "regex_best_effort"}})
            if len(chunks) > MAX_CHUNKS:
                raise ValueError(f"Corpus exceeds {MAX_CHUNKS} chunks; previous index retained")
    write_json(Path(work) / "chunks.json", chunks)
    return work


def publish_index(work, embedder=None):
    import faiss
    import numpy as np
    discovery = read_json(Path(work) / "discovery.json")
    chunks = read_json(Path(work) / "chunks.json")
    # Do not publish a mixed snapshot when a source changed between tasks.
    for item in discovery["files"]:
        if hashlib.sha256(Path(item["path"]).read_bytes()).hexdigest() != item["sha256"]:
            raise ValueError(f"Document changed before publication: {item['source']}")
    if chunks and embedder is None:
        from sentence_transformers import SentenceTransformer
        embedder = SentenceTransformer(discovery["model"], device="cpu")
    dimension = int(embedder.get_sentence_embedding_dimension()) if chunks else 384
    index = faiss.IndexFlatL2(dimension)
    for offset in range(0, len(chunks), 64):
        vectors = np.asarray(embedder.encode([c["text"] for c in chunks[offset:offset + 64]],
                             normalize_embeddings=True, show_progress_bar=False), dtype="float32")
        if vectors.shape != (len(chunks[offset:offset + 64]), dimension) or not np.isfinite(vectors).all():
            raise ValueError("Embedding model returned invalid vectors")
        index.add(vectors)
    root = index_directory()
    version = uuid.uuid4().hex
    snapshot = root / "versions" / version
    snapshot.mkdir(parents=True)
    faiss.write_index(index, str(snapshot / "index.faiss"))
    write_json(snapshot / "chunks.json", chunks)
    manifest = {"version": version, "fingerprint": discovery["fingerprint"],
                "model": discovery["model"], "dimension": dimension,
                "chunk_count": len(chunks), "document_count": len(discovery["files"]),
                "normalized": True, "metric": "l2", "pipeline_version": PIPELINE_VERSION,
                "published_at": datetime.now(timezone.utc).isoformat()}
    write_json(snapshot / "manifest.json", manifest)
    # The successful fingerprint and aligned artifacts become visible together.
    write_json(root / "current.json", manifest)
    write_json(Path(work) / "report.json", manifest)
    return manifest
